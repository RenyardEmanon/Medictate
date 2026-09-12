import os
import uuid

from flask import current_app
from werkzeug.utils import secure_filename
from docx import Document

from stt import transcribe_audio, TranscriptionError
from llm import refine_transcript, RefinementError


def allowed_audio(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in current_app.config["ALLOWED_AUDIO_EXTENSIONS"]


def build_docx(report):
    """Render the report's transcript text into a formatted Word document."""
    doc = Document()
    doc.add_heading("Radiology Report", level=0)

    patient = report.patient
    doctor = report.author
    scan = report.scan

    doc.add_paragraph(f"Patient: {patient.name}")
    doc.add_paragraph(f"Reporting Doctor: {doctor.name} ({doctor.specialty or 'Radiologist'})")
    if scan:
        doc.add_paragraph(f"Scan: {scan.filename}  |  Body part: {scan.body_part or 'N/A'}")
    doc.add_paragraph(f"Report date: {report.updated_at.strftime('%Y-%m-%d %H:%M UTC')}")
    doc.add_paragraph("")

    doc.add_heading("Findings & Impression", level=1)
    doc.add_paragraph(report.transcript_text or "")

    reports_dir = current_app.config["REPORTS_FOLDER"]
    os.makedirs(reports_dir, exist_ok=True)
    filename = f"report_{report.id}_{uuid.uuid4().hex[:8]}.docx"
    path = os.path.join(reports_dir, filename)
    doc.save(path)
    return path


def run_transcription_pipeline(audio_file):
    """Validate audio, run ASR + clinical cleanup, return final text.

    Raises:
        ValueError - for bad/missing/empty audio input
        TranscriptionError - for ASR provider failures (has .status_code)
    """
    if audio_file.filename == "" or not allowed_audio(audio_file.filename):
        raise ValueError("Unsupported or missing audio file")

    audio_bytes = audio_file.read()
    if not audio_bytes:
        raise ValueError("Uploaded audio file is empty")

    safe_name = secure_filename(audio_file.filename)

    tmp_path = os.path.join(current_app.config["UPLOAD_FOLDER"], f"audio_{uuid.uuid4().hex}_{safe_name}")
    with open(tmp_path, "wb") as f:
        f.write(audio_bytes)

    asr_result = transcribe_audio(audio_bytes, safe_name)
    raw_text = asr_result["transcript"]

    try:
        cleaned = refine_transcript(raw_text)
    except RefinementError as e:
        current_app.logger.warning("Clinical cleanup failed, using raw transcript: %s", e.message)
        cleaned = raw_text

    final_text = cleaned if cleaned is not None else raw_text
    if not final_text.strip():
        raise ValueError("No clinical content was detected in the recording. Please try again.")

    return final_text