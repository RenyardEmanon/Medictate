import os
import uuid

from flask import Blueprint, request, jsonify, current_app, send_from_directory
from flask_login import login_required, current_user
from flask_mail import Message
from werkzeug.utils import secure_filename
from docx import Document

from extensions import db, mail
from models import Scan, Report, User
from stt import transcribe_audio, TranscriptionError
from llm import refine_transcript, RefinementError

reports_bp = Blueprint("reports", __name__, url_prefix="/api/reports")


def _allowed_audio(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in current_app.config["ALLOWED_AUDIO_EXTENSIONS"]


def build_docx(report: Report):
    """Render the report's transcript text into a formatted Word document."""
    doc = Document()
    doc.add_heading("Radiology Report", level=0)

    patient = report.patient
    doctor = report.author
    scan = report.scan  # may be None now

    doc.add_paragraph(f"Patient: {patient.name}")
    doc.add_paragraph(f"Reporting Doctor: {doctor.name} ({doctor.specialty or 'Radiologist'})")
    if scan:
        doc.add_paragraph(f"Scan: {scan.filename}  |  Body part: {scan.body_part or 'N/A'}")
    doc.add_paragraph(f"Report date: {report.updated_at.strftime('%Y-%m-%d %H:%M UTC')}")
    doc.add_paragraph("")

    doc.add_heading("Findings & Impression", level=1)
    doc.add_paragraph(report.transcript_text or "")

    reports_dir = current_app.config["REPORTS_FOLDER"]
    os.makedirs(reports_dir, exist_ok=True)
    filename = f"report_{report.id}_{uuid.uuid4().hex[:8]}.docx"
    path = os.path.join(reports_dir, filename)
    doc.save(path)
    return path


def _run_transcription_pipeline(audio_file):
    """Shared by both dictation routes: validate audio, run ASR + cleanup, return final text."""
    if audio_file.filename == "" or not _allowed_audio(audio_file.filename):
        raise ValueError("Unsupported or missing audio file")

    audio_bytes = audio_file.read()
    if not audio_bytes:
        raise ValueError("Uploaded audio file is empty")

    safe_name = secure_filename(audio_file.filename)

    # Keep a copy of the raw audio for reference/debugging
    tmp_path = os.path.join(current_app.config["UPLOAD_FOLDER"], f"audio_{uuid.uuid4().hex}_{safe_name}")
    with open(tmp_path, "wb") as f:
        f.write(audio_bytes)

    asr_result = transcribe_audio(audio_bytes, safe_name)  # raises TranscriptionError
    raw_text = asr_result["transcript"]

    try:
        cleaned = refine_transcript(raw_text)
    except RefinementError as e:
        current_app.logger.warning("Clinical cleanup failed, using raw transcript: %s", e.message)
        cleaned = raw_text

    final_text = cleaned if cleaned is not None else raw_text
    if not final_text.strip():
        raise ValueError("No clinical content was detected in the recording. Please try again.")

    return final_text


@reports_bp.post("/record")
@login_required
def record_standalone():
    """Doctor dictates a report for a patient directly - no scan required."""
    if current_user.role != "doctor":
        return jsonify({"error": "Only doctors can dictate reports"}), 403

    patient_id = request.form.get("patient_id", type=int)
    patient = User.query.get(patient_id) if patient_id else None
    if not patient or patient.role != "patient":
        return jsonify({"error": "Valid patient_id is required"}), 400
    if patient.assigned_doctor_id != current_user.id:
        return jsonify({"error": "This patient is not assigned to you"}), 403

    if "audio" not in request.files:
        return jsonify({"error": "No audio file part named 'audio'"}), 400

    try:
        final_text = _run_transcription_pipeline(request.files["audio"])
    except (ValueError, TranscriptionError) as e:
        message = e.message if isinstance(e, TranscriptionError) else str(e)
        status = e.status_code if isinstance(e, TranscriptionError) else 400
        return jsonify({"error": message}), status

    report = Report(patient_id=patient.id, author_id=current_user.id, transcript_text=final_text)
    db.session.add(report)
    db.session.flush()
    report.docx_path = build_docx(report)
    db.session.commit()

    return jsonify({"message": "Transcribed", "report": report.to_dict()})


@reports_bp.post("/<int:scan_id>/transcribe")
@login_required
def transcribe(scan_id):
    """Doctor dictates a report tied to a specific scan."""
    if current_user.role != "doctor":
        return jsonify({"error": "Only doctors can dictate reports"}), 403

    scan = Scan.query.get_or_404(scan_id)
    if scan.doctor_id != current_user.id:
        return jsonify({"error": "You are not the reporting doctor for this scan"}), 403

    if "audio" not in request.files:
        return jsonify({"error": "No audio file part named 'audio'"}), 400

    try:
        final_text = _run_transcription_pipeline(request.files["audio"])
    except (ValueError, TranscriptionError) as e:
        message = e.message if isinstance(e, TranscriptionError) else str(e)
        status = e.status_code if isinstance(e, TranscriptionError) else 400
        return jsonify({"error": message}), status

    report = Report.query.filter_by(scan_id=scan.id).first()
    if not report:
        report = Report(scan_id=scan.id, patient_id=scan.patient_id, author_id=current_user.id)
        db.session.add(report)

    report.transcript_text = final_text
    db.session.flush()
    report.docx_path = build_docx(report)
    scan.status = "reported"
    db.session.commit()

    return jsonify({"message": "Transcribed", "report": report.to_dict()})


@reports_bp.put("/<int:report_id>/correct")
@login_required
def correct(report_id):
    report = Report.query.get_or_404(report_id)
    if current_user.id != report.author_id:
        return jsonify({"error": "Only the authoring doctor can edit this report"}), 403

    data = request.get_json(force=True) or {}
    corrected_text = data.get("transcript_text")
    if corrected_text is None:
        return jsonify({"error": "transcript_text is required"}), 400

    report.transcript_text = corrected_text
    report.docx_path = build_docx(report)
    db.session.commit()

    return jsonify({"message": "Report updated", "report": report.to_dict()})


@reports_bp.get("/<int:report_id>/download")
@login_required
def download_own(report_id):
    report = Report.query.get_or_404(report_id)
    if current_user.id not in (report.author_id, report.patient_id):
        return jsonify({"error": "Not authorized"}), 403
    directory, name = os.path.split(report.docx_path)
    return send_from_directory(directory, name, as_attachment=True, download_name="radiology_report.docx")


@reports_bp.post("/<int:report_id>/send-to-patient")
@login_required
def send_to_patient(report_id):
    report = Report.query.get_or_404(report_id)
    if current_user.id != report.author_id:
        return jsonify({"error": "Only the authoring doctor can send this report"}), 403
    if not report.docx_path:
        return jsonify({"error": "Report has no generated document yet"}), 400

    api_key = report.generate_api_key()
    report.sent_to_patient = True
    db.session.commit()

    download_link = f"{current_app.config['PUBLIC_BASE_URL']}/api/reports/public-download/{report.id}?key={api_key}"

    try:
        msg = Message(
            subject="Your Radiology Report is Ready",
            recipients=[report.patient.email],
            body=(
                f"Hello {report.patient.name},\n\n"
                f"Your radiology report from Dr. {report.author.name} is ready.\n"
                f"You can securely download it here:\n{download_link}\n\n"
                "This link is unique to you - please don't share it.\n\n"
                "Regards,\nRadiology Team"
            ),
        )
        mail.send(msg)
    except Exception as exc:
        current_app.logger.warning("Email send failed: %s", exc)
        return jsonify({
            "message": "API key generated, but email failed to send (check MAIL_* config).",
            "download_link": download_link,
        }), 202

    return jsonify({"message": f"Report emailed to {report.patient.email}", "download_link": download_link})


@reports_bp.get("/public-download/<int:report_id>")
def public_download(report_id):
    report = Report.query.get_or_404(report_id)
    supplied_key = request.args.get("key")
    if not supplied_key or supplied_key != report.api_key:
        return jsonify({"error": "Invalid or missing API key"}), 401
    directory, name = os.path.split(report.docx_path)
    return send_from_directory(directory, name, as_attachment=True, download_name="radiology_report.docx")
